import importlib.util
import json
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest

from web.agent_runtime import (
    AgentTemplateError,
    compile_agent_template,
    interrupt_python_run,
    is_python_run_active,
    migrate_legacy_agent_template,
    run_agent_python,
    run_script_python,
    stream_agent_python,
    stream_script_python,
)


PARAMETERS = {
    "model": 'model-"quoted"',
    "model_provider": "provider",
    "api_key": "secret-key",
    "base_url": "https://example.test/v1",
    "system_prompt": "",
    "human_message": "line 1\nline 2",
}


def _run_stream_worker(code, runtime_label="Script"):
    payload = json.dumps(
        {"code": code, "runtime_label": runtime_label, "stream": True},
        ensure_ascii=False,
    ).encode("utf-8")
    process = subprocess.run(
        [sys.executable, "-m", "web.agent_worker"],
        input=payload,
        capture_output=True,
        check=False,
        timeout=10,
    )
    events = []
    for raw_line in process.stdout.decode("utf-8").split("\n"):
        if not raw_line:
            continue
        assert raw_line.startswith("\x1e")
        events.append(json.loads(raw_line[1:]))
    return process, events


def test_compile_agent_template_uses_python_literals_and_none():
    compiled = compile_agent_template(
        "values = [${model}, ${system_prompt}, ${human_message}]",
        PARAMETERS,
    )

    namespace = {}
    exec(compiled, namespace)
    assert namespace["values"] == ['model-"quoted"', None, "line 1\nline 2"]


def test_literal_values_override_fields_and_commented_placeholders_do_not_execute():
    compiled = compile_agent_template(
        'model = "hard-coded"\n# ignored = ${model}\nresponse = model',
        PARAMETERS,
    )

    namespace = {}
    exec(compiled, namespace)
    assert namespace["response"] == "hard-coded"
    assert "# ignored = 'model-\"quoted\"'" in compiled


def test_compile_agent_template_rejects_unknown_placeholder():
    with pytest.raises(AgentTemplateError, match="unknown_value"):
        compile_agent_template("response = ${unknown_value}", PARAMETERS)


def test_compile_agent_template_rejects_malformed_placeholder():
    with pytest.raises(AgentTemplateError, match="无法识别"):
        compile_agent_template("response = ${model", PARAMETERS)


def test_compile_agent_template_allows_nested_python_dicts():
    compiled = compile_agent_template(
        'response = {"model": ${model}, "extra_body": {"thinking": {"type": "disabled"}}}',
        PARAMETERS,
    )

    namespace = {}
    exec(compiled, namespace)
    assert namespace["response"] == {
        "model": 'model-"quoted"',
        "extra_body": {"thinking": {"type": "disabled"}},
    }


def test_compile_agent_template_allows_sets_and_nested_json():
    compiled = compile_agent_template(
        'import json\nresponse = {"values": {"one", "two"}, "payload": '
        'json.loads(\'{"outer": {"items": [{"enabled": true}]}}\'), "model": ${model}}',
        PARAMETERS,
    )

    namespace = {}
    exec(compiled, namespace)
    assert namespace["response"]["values"] == {"one", "two"}
    assert namespace["response"]["payload"] == {
        "outer": {"items": [{"enabled": True}]}
    }
    assert namespace["response"]["model"] == 'model-"quoted"'


def test_migrate_legacy_agent_template_only_changes_placeholders():
    legacy = (
        'response = {"model": {{model}}, '
        '"extra_body": {"thinking": {"type": "disabled"}}}'
    )

    assert migrate_legacy_agent_template(legacy) == (
        'response = {"model": ${model}, '
        '"extra_body": {"thinking": {"type": "disabled"}}}'
    )


def test_worker_supports_imports_and_only_returns_user_print_output():
    result = run_agent_python(
        "import math\nprint('sqrt', math.sqrt(81))\nresponse = {'ok': True}",
        PARAMETERS,
    )

    assert result["ok"] is True
    assert "sqrt 9.0" in result["logs"]
    assert "'ok': True" not in result["logs"]


def test_worker_allows_user_editable_rich_print():
    result = run_agent_python(
        "from rich import print\nresponse = {'ok': True}\nprint(response)",
        PARAMETERS,
    )

    assert result["ok"] is True
    assert "'ok': True" in result["logs"]


def test_worker_does_not_inherit_delayed_annotations_from_worker_module():
    result = run_agent_python(
        """from langchain.agents.structured_output import ToolStrategy
from pydantic import BaseModel

class NobelWinner(BaseModel):
    name: str

class NobelAnswer(BaseModel):
    winners: list[NobelWinner]

response = ToolStrategy(NobelAnswer)
print(type(response).__name__)
""",
        PARAMETERS,
    )

    assert result["ok"] is True
    assert "ToolStrategy" in result["logs"]


def test_worker_success_can_have_empty_logs():
    result = run_agent_python("response = {'ok': True}", PARAMETERS)

    assert result == {"ok": True, "logs": "", "response": {"ok": True}}


def test_stream_worker_emits_unicode_lines_flush_and_pydantic_response():
    process, events = _run_stream_worker(
        """import sys
from pydantic import BaseModel

class Result(BaseModel):
    message: str

print('第一行 🍖')
sys.stdout.write('无换行输出')
sys.stdout.flush()
response = Result(message='完成 🍖')
"""
    )

    assert process.returncode == 0
    assert events == [
        {"type": "log", "text": "第一行 🍖\n"},
        {"type": "log", "text": "无换行输出"},
        {
            "type": "result",
            "result": {
                "ok": True,
                "logs": "",
                "response": {"message": "完成 🍖"},
            },
        },
    ]


def test_stream_worker_emits_traceback_before_failed_result():
    process, events = _run_stream_worker("raise ValueError('stream failure')")

    assert process.returncode == 0
    logs = "".join(
        event["text"] for event in events if event["type"] == "log"
    )
    assert "Script 执行失败" in logs
    assert "ValueError: stream failure" in logs
    assert "完整 Traceback" in logs
    assert events[-1] == {
        "type": "result",
        "result": {"ok": False, "logs": ""},
    }


def test_stream_runtime_delivers_logs_before_structured_result():
    logs = []

    result = stream_script_python(
        "print('line one')\nprint('第二行')\nresponse = {'answer': 42}",
        logs.append,
        "stream-runtime-success",
    )

    assert logs == ["line one\n", "第二行\n"]
    assert result == {
        "ok": True,
        "logs": "",
        "response": {"answer": 42},
    }


def test_stream_agent_compiles_parameters_before_emitting():
    logs = []

    result = stream_agent_python(
        "print(${model})\nresponse = ${human_message}",
        PARAMETERS,
        logs.append,
        "stream-agent-success",
    )

    assert logs == ['model-"quoted"\n']
    assert result["response"] == "line 1\nline 2"


def test_stream_runtime_interrupts_after_delivering_existing_lines():
    logs = []
    run_id = "stream-runtime-interrupt"

    with ThreadPoolExecutor(max_workers=1) as executor:
        future = executor.submit(
            stream_script_python,
            "import time\nprint('started', flush=True)\ntime.sleep(30)",
            logs.append,
            run_id,
            30,
        )
        deadline = time.monotonic() + 5
        while not logs and time.monotonic() < deadline:
            time.sleep(0.02)

        assert logs == ["started\n"]
        assert interrupt_python_run(run_id) is True
        result = future.result(timeout=5)

    assert result == {"ok": False, "interrupted": True, "logs": ""}


def test_worker_success_does_not_require_response():
    result = run_agent_python("print('普通 Python 代码执行完成')", PARAMETERS)

    assert result == {"ok": True, "logs": "普通 Python 代码执行完成\n"}


def test_worker_returns_json_structured_response():
    result = run_agent_python(
        """import json
from pathlib import Path
from pydantic import BaseModel

class Item(BaseModel):
    name: str
    count: int

response = {
    "item": Item(name="demo", count=2),
    "payload": json.loads('{"enabled": true}'),
    "path": Path("outputs/result.json"),
}
""",
        PARAMETERS,
    )

    assert result["ok"] is True
    assert result["logs"] == ""
    assert result["response"]["item"] == {"name": "demo", "count": 2}
    assert result["response"]["payload"] == {"enabled": True}
    assert Path(result["response"]["path"]) == Path("outputs/result.json")


def test_worker_executes_installed_langchain_tools_and_middleware():
    result = run_agent_python(
        """from langchain.agents.middleware import AgentMiddleware, ToolRetryMiddleware
from langchain_core.tools import tool
from langchain_tavily import TavilySearch

@tool
def double(value: int) -> int:
    \"\"\"把输入整数乘以二。\"\"\"
    return value * 2

class AuditMiddleware(AgentMiddleware):
    def wrap_tool_call(self, request, handler):
        return handler(request)

middlewares = [
    ToolRetryMiddleware(max_retries=1, tools=[double]),
    AuditMiddleware(),
]
response = {
    "builtin_tool": TavilySearch.__name__,
    "custom_tool": double.name,
    "tool_result": double.invoke({"value": 21}),
    "middleware": [type(item).__name__ for item in middlewares],
}
""",
        PARAMETERS,
    )

    assert result["ok"] is True
    assert result["response"] == {
        "builtin_tool": "TavilySearch",
        "custom_tool": "double",
        "tool_result": 42,
        "middleware": ["ToolRetryMiddleware", "AuditMiddleware"],
    }


def test_script_worker_executes_langchain_pydantic_stdlib_and_third_party_packages():
    result = run_script_python(
        """import statistics
from langchain_core.tools import tool
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

@tool
def triple(value: int) -> int:
    \"\"\"把输入整数乘以三。\"\"\"
    return value * 3

class ScriptResult(BaseModel):
    value: int
    column: str
    mean: float

print("script imports ready")
response = ScriptResult(
    value=triple.invoke({"value": 14}),
    column=get_column_letter(28),
    mean=statistics.mean([2, 4, 6]),
)
"""
    )

    assert result["ok"] is True
    assert result["logs"] == "script imports ready\n"
    assert result["response"] == {"value": 42, "column": "AB", "mean": 4.0}


def test_agent_worker_initializes_custom_anthropic_client():
    result = run_agent_python(
        '''import httpx
from anthropic import Anthropic
from langchain.chat_models import init_chat_model
from langchain.agents import create_agent
from langchain_core.messages import HumanMessage

http_client = httpx.Client(trust_env=False)

model = init_chat_model(
    model="custom-anthropic-model",
    model_provider="anthropic",
    api_key="test-key",
    base_url="https://example.test",
)

model._client = Anthropic(
    api_key="test-key",
    base_url="https://example.test",
    http_client=http_client,
)
''',
        {},
    )

    assert result == {"ok": True, "logs": ""}


def test_script_worker_does_not_replace_agent_placeholders():
    result = run_script_python('response = "${model}"')

    assert result == {"ok": True, "logs": "", "response": "${model}"}


def test_script_worker_times_out_and_uses_script_error_label():
    timed_out = run_script_python("import time\ntime.sleep(5)", timeout_seconds=0.2)
    failed = run_script_python("raise ValueError('script failure')")

    assert timed_out["ok"] is False
    assert "Script 执行超时" in timed_out["logs"]
    assert failed["ok"] is False
    assert "Script 执行失败" in failed["logs"]
    assert "ValueError: script failure" in failed["logs"]


def test_worker_reports_missing_dependency_without_installing():
    assert importlib.util.find_spec("pendulum") is None

    result = run_agent_python(
        """from pendulum import now

response = now("Asia/Shanghai").to_iso8601_string()
""",
        {},
    )

    assert result["ok"] is False
    assert "缺少 Python 模块: pendulum" in result["logs"]
    assert "系统不会自动安装依赖" in result["logs"]
    assert "加入 pyproject.toml" in result["logs"]
    assert "执行 uv sync 后重试" in result["logs"]
    assert "No module named 'pendulum'" in result["logs"]
    assert "ModuleNotFoundError" in result["logs"]
    assert "Traceback" in result["logs"]
    assert "response" not in result
    assert importlib.util.find_spec("pendulum") is None


def test_worker_returns_user_exception_and_traceback():
    result = run_agent_python(
        "print('before failure')\nraise ValueError('boom')",
        PARAMETERS,
    )

    assert result["ok"] is False
    assert "before failure" in result["logs"]
    assert "ValueError: boom" in result["logs"]
    assert "Traceback" in result["logs"]


def test_worker_times_out_and_terminates_process():
    result = run_agent_python(
        "import time\ntime.sleep(5)\nresponse = None",
        PARAMETERS,
        timeout_seconds=0.2,
    )

    assert result["ok"] is False
    assert "执行超时" in result["logs"]


def test_worker_can_be_interrupted_by_run_id_without_returning_partial_logs(tmp_path):
    started_file = tmp_path / "started.json"
    run_id = "interrupt-runtime-test"
    code = (
        "import json, os, time\n"
        f"open({str(started_file)!r}, 'w', encoding='utf-8').write("
        "json.dumps({'pid': os.getpid()}))\n"
        "print('partial output')\n"
        "time.sleep(30)\n"
        "response = {'finished': True}\n"
    )

    with ThreadPoolExecutor(max_workers=1) as executor:
        future = executor.submit(run_script_python, code, 30, run_id)
        deadline = time.monotonic() + 5
        while not started_file.exists() and time.monotonic() < deadline:
            time.sleep(0.02)

        assert started_file.exists()
        assert is_python_run_active(run_id) is True
        assert interrupt_python_run(run_id) is True
        result = future.result(timeout=5)

    assert result == {"ok": False, "interrupted": True, "logs": ""}
    assert is_python_run_active(run_id) is False
    assert json.loads(started_file.read_text(encoding="utf-8"))["pid"] > 0


def test_interrupt_before_worker_start_prevents_execution(tmp_path):
    marker = tmp_path / "should-not-exist.txt"
    run_id = "pre-cancel-runtime-test"

    assert interrupt_python_run(run_id) is False
    result = run_script_python(
        f"open({str(marker)!r}, 'w').write('ran')",
        run_id=run_id,
    )

    assert result == {"ok": False, "interrupted": True, "logs": ""}
    assert marker.exists() is False


def test_interrupt_terminates_worker_child_process_tree(tmp_path):
    started_file = tmp_path / "child-started.txt"
    survived_file = tmp_path / "child-survived.txt"
    run_id = "interrupt-process-tree-test"
    child_code = (
        "import time\n"
        "time.sleep(1.5)\n"
        f"open({str(survived_file)!r}, 'w').write('survived')\n"
    )
    code = (
        "import subprocess, sys, time\n"
        f"child = subprocess.Popen([sys.executable, '-c', {child_code!r}])\n"
        f"open({str(started_file)!r}, 'w').write(str(child.pid))\n"
        "time.sleep(30)\n"
    )

    with ThreadPoolExecutor(max_workers=1) as executor:
        future = executor.submit(run_script_python, code, 30, run_id)
        deadline = time.monotonic() + 5
        while not started_file.exists() and time.monotonic() < deadline:
            time.sleep(0.02)

        assert started_file.exists()
        assert interrupt_python_run(run_id) is True
        result = future.result(timeout=5)

    time.sleep(1.8)
    assert result == {"ok": False, "interrupted": True, "logs": ""}
    assert survived_file.exists() is False


def test_worker_returns_logs_without_redaction():
    parameters = {**PARAMETERS, "api_key": "a"}
    result = run_agent_python(
        "print('Agent Pydantic call')\nprint(${api_key})\nraise RuntimeError('fatal')",
        parameters,
    )

    assert result["ok"] is False
    assert "Agent Pydantic call" in result["logs"]
    assert "RuntimeError: fatal" in result["logs"]
    assert "***" not in result["logs"]


def test_worker_protocol_preserves_unicode_outside_gbk():
    result = run_agent_python(
        "print('中文日志 🍖')\nresponse = {'message': '执行成功 🍖'}\nprint(response)",
        PARAMETERS,
    )

    assert result["ok"] is True
    assert "中文日志 🍖" in result["logs"]
    assert "执行成功 🍖" in result["logs"]
